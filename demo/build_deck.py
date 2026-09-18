"""Build the presentation from slides.md, the script, and committed research metrics."""

import csv
import hashlib
import json
import posixpath
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
DEMO = ROOT / "demo"
BG = "0B1220"
PANEL = "142137"
WHITE = "F4F7FC"
MUTED = "B1C0D4"
MINT = "6DE5C1"
VIOLET = "B9A3FF"
AMBER = "F5C978"
RED = "FF929D"
BLUE = "90C7F6"
LINE = "34465F"
FONT = "Arial"
WIDTH = 13.333333
HEIGHT = 7.5


def color(value):
    return RGBColor.from_string(value)


def box(slide, x, y, w, h, fill=PANEL, stroke=None, rounded=True):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        shape.adjustments[0] = 0.12
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    if stroke:
        shape.line.color.rgb = color(stroke)
        shape.line.width = Pt(1.25)
    else:
        shape.line.fill.background()
    return shape


def text(slide, value, x, y, w, h, size=24, ink=WHITE, bold=False, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for index, line in enumerate(value.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.alignment = align
        paragraph.space_before = paragraph.space_after = Pt(0)
        paragraph.line_spacing = 1.05
        paragraph.font.name = FONT
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color(ink)
    return shape


def arrow(slide, x1, y1, x2, y2, ink=MINT, head=True):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    line.line.color.rgb = color(ink)
    line.line.width = Pt(1.7)
    if head:
        end = OxmlElement("a:tailEnd")
        end.set("type", "triangle")
        line.line._get_or_add_ln().append(end)
    return line


def pill(slide, value, x, y, w, ink=MINT, size=16):
    box(slide, x, y, w, 0.43, fill=PANEL, stroke=ink)
    text(slide, value, x + 0.09, y + 0.01, w - 0.18, 0.4, size, ink, True, PP_ALIGN.CENTER)


def node(slide, value, x, y, w, h=0.75, ink=MINT):
    box(slide, x, y, w, h, stroke=ink)
    text(slide, value, x + 0.08, y + 0.02, w - 0.16, h - 0.04, 21, ink, True, PP_ALIGN.CENTER)


def source():
    raw = (DEMO / "slides.md").read_text()
    pattern = r"^## Slide (\d+) — ([^\n]+)\n(.*?)(?=^## Slide |^## Source-of-truth|\Z)"
    result = []
    for number, title, body in re.findall(pattern, raw, re.M | re.S):
        copy = re.search(r"### On-slide text\n```text\n(.*?)\n```", body, re.S)
        if not copy:
            raise ValueError(f"Missing slide text: {number}")
        result.append((int(number), title, copy[1].splitlines(), body))
    if [entry[0] for entry in result] != list(range(1, 12)):
        raise ValueError("Expected exactly eleven ordered slide specifications")
    return result


def new_slide(prs, number, title, section, note):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color(BG)
    text(slide, "AQL / " + section, 0.55, 0.28, 11.8, 0.3, 13, MINT, True)
    if number not in (1, 11):
        text(slide, title, 0.55, 0.83, 12.2, 1.12, 34, bold=True)
    text(slide, "AGENTIC QUANT LAB", 0.55, 7.04, 7, 0.22, 11, MUTED)
    text(slide, f"{number:02} / 11", 11.7, 7.01, 1.08, 0.25, 12, MUTED, align=PP_ALIGN.RIGHT)
    for index in range(11):
        box(
            slide,
            0.55 + index * 1.115,
            7.37,
            1.045,
            0.025,
            fill=MINT if index < number else LINE,
            rounded=False,
        )
    slide.notes_slide.notes_text_frame.text = note
    return slide


def title_slide(s, t):
    text(s, "Agentic\nQuant Lab", 0.55, 1.25, 7.6, 2.0, 62, bold=True)
    text(s, "\n".join(t[:2]), 0.6, 3.52, 7.0, 1.1, 26, MUTED)
    text(s, "\n".join(t[2:4]), 0.6, 5.25, 7.0, 0.94, 28, MINT, True)
    for i, value in enumerate(t[4:8]):
        y = 1.5 + i * 1.14
        node(s, value, 8.55, y, 3.7, 0.83, VIOLET if i == 0 else MINT)
        if i < 3:
            arrow(s, 10.4, y + 0.85, 10.4, y + 1.1)
    text(s, t[8], 8.35, 6.32, 4.1, 0.32, 14, MUTED, align=PP_ALIGN.CENTER)


def problem_slide(s, t):
    box(s, 0.55, 2.05, 12.23, 1.27, fill="251D2C")
    text(s, t[0], 0.8, 2.17, 3.35, 0.76, 18, RED, True)
    for i, value in enumerate(t[1:4]):
        x = 4.3 + i * 2.78
        node(s, value, x, 2.31, 2.25, 0.65, RED)
        if i < 2:
            arrow(s, x + 2.3, 2.64, x + 2.68, 2.64, RED)
    box(s, 0.55, 3.55, 12.23, 1.3)
    text(s, t[4], 0.8, 3.78, 2.32, 0.7, 18, MINT, True)
    for i, value in enumerate(t[5:10]):
        x = 3.3 + i * 1.88
        node(s, value, x, 3.88, 1.64, 0.62)
        if i < 4:
            arrow(s, x + 1.65, 4.19, x + 1.83, 4.19)
    for i, value in enumerate(t[10:14]):
        text(s, value, 0.6 + i * 3.1, 5.09, 2.9, 0.55, 17, MUTED)
    text(s, t[14], 0.55, 5.98, 12.23, 0.55, 31, WHITE, True)
    text(s, t[15], 0.55, 6.65, 12.0, 0.25, 14, MUTED)


def architecture_slide(s, t):
    pill(s, t[0], 0.55, 2.0, 2.5, VIOLET)
    pill(s, t[1], 3.5, 2.0, 5.25)
    node(s, t[2], 0.55, 2.7, 2.5, ink=VIOLET)
    node(s, t[3], 0.55, 3.93, 2.5, ink=VIOLET)
    arrow(s, 1.8, 3.48, 1.8, 3.89, VIOLET)
    arrow(s, 3.06, 4.3, 3.28, 4.3, VIOLET, False)
    arrow(s, 3.28, 4.3, 3.28, 3.08, VIOLET, False)
    arrow(s, 3.28, 3.08, 3.49, 3.08, VIOLET)
    xs = [3.5, 5.85, 8.2, 10.55]
    for i, x in enumerate(xs):
        node(s, t[4 + i], x, 2.7, 2.18)
        if i < 3:
            arrow(s, x + 2.2, 3.08, xs[i + 1] - 0.02, 3.08)
    arrow(s, 11.64, 3.48, 11.64, 4.17)
    for i, x in enumerate(reversed(xs)):
        value = t[8 + i]
        node(s, value, x, 4.2, 2.18, ink=MUTED if i == 3 else MINT)
        if i < 3:
            arrow(s, x - 0.01, 4.58, x - 0.16, 4.58, MUTED if i == 2 else MINT)
    text(s, t[12], 3.5, 5.04, 2.18, 0.3, 14, MUTED, align=PP_ALIGN.CENTER)
    box(s, 0.55, 5.68, 12.18, 0.86, fill="182435", stroke=LINE)
    text(s, t[13], 0.8, 5.84, 3.45, 0.46, 19, BLUE, True)
    text(s, t[14], 4.45, 5.86, 7.98, 0.44, 21, MUTED)
    arrow(s, 1.8, 4.72, 1.8, 5.65, BLUE)
    arrow(s, 7.83, 3.47, 8.11, 3.7, BLUE, False)
    arrow(s, 8.11, 3.7, 8.11, 5.65, BLUE)
    arrow(s, 9.29, 4.98, 9.29, 5.65, BLUE)
    text(s, t[15], 0.55, 6.66, 12.2, 0.27, 14, MUTED)


def controls_slide(s, t):
    for i, (symbol, ink) in enumerate([("0", BLUE), ("+", MINT), ("t+1", AMBER)]):
        x = 0.55 + i * 4.15
        box(s, x, 2.14, 3.93, 3.05, stroke=LINE)
        text(s, t[i * 3], x + 0.24, 2.39, 3.45, 0.4, 20, ink, True)
        text(s, symbol, x + 0.24, 2.99, 3.45, 0.88, 56, ink, True)
        text(s, "\n".join(t[i * 3 + 1 : i * 3 + 3]), x + 0.24, 4.01, 3.45, 0.87, 26)
    for i, value in enumerate(t[9:13]):
        text(s, value, 0.55 + i * 3.1, 5.46, 2.93, 0.4, 20, MUTED)
    text(s, t[13], 0.55, 6.04, 12.23, 0.5, 29, MINT, True)
    text(s, t[14], 0.55, 6.65, 12.23, 0.27, 14, MUTED)


def research_slide(s, t):
    pill(s, t[0], 10.0, 0.25, 2.77, AMBER, 14)
    for i in range(5):
        name, status, reason = t[1 + i * 3 : 4 + i * 3]
        y = 2.05 + i * 0.76
        ink = AMBER if i in (0, 3) else BLUE
        box(s, 0.55, y, 12.23, 0.65)
        box(s, 0.55, y, 0.055, 0.65, fill=ink, rounded=False)
        text(s, name, 0.77, y + 0.07, 3.45, 0.49, 21, bold=True)
        text(s, status, 4.3, y + 0.07, 2.85, 0.49, 18, ink, True)
        text(s, reason, 7.22, y + 0.07, 5.31, 0.49, 19, MUTED)
    text(s, t[16], 0.55, 6.01, 12.23, 0.51, 29, MINT, True)
    text(s, t[17], 0.55, 6.66, 12.23, 0.27, 14, MUTED)


def trend_slide(s, t):
    text(s, t[0], 0.55, 1.98, 12, 0.29, 15, AMBER, True)
    text(s, t[1], 0.55, 2.35, 12, 0.3, 16, MUTED)
    text(s, t[2], 3.12, 2.86, 3.5, 0.36, 19, MINT, True)
    text(s, t[3], 8.2, 2.86, 4.2, 0.36, 19, BLUE, True)
    with (ROOT / "research/etf_trend/results/metrics.csv").open() as handle:
        rows = {
            row["model"]: row
            for row in csv.DictReader(handle)
            if row["defense"] == "BIL" and row["one_way_cost_bps"] == "20"
        }
    models = [
        ("SPY", "SPY"),
        ("STATIC80", "Static 80/20"),
        ("STATIC70", "Static 70/30"),
        ("ABS12", "Absolute12"),
        ("SMA10", "SMA10"),
    ]
    for i, (model, label) in enumerate(models):
        y = 3.4 + i * 0.43
        cagr = float(rows[model]["cagr"]) * 100
        drawdown = float(rows[model]["max_drawdown"]) * 100
        ink = AMBER if model in ("ABS12", "SMA10") else MINT
        text(s, label, 0.55, y - 0.025, 2.45, 0.34, 20, bold=True)
        box(s, 3.12, y + 0.015, cagr / 16 * 3.05, 0.22, fill=ink, rounded=False)
        text(s, f"{cagr:.2f}%", 6.3, y - 0.025, 1.38, 0.34, 20, ink, True)
        box(s, 8.2, y + 0.015, abs(drawdown) / 40 * 2.68, 0.22, fill=BLUE, rounded=False)
        text(s, f"{drawdown:.2f}%", 11.08, y - 0.025, 1.68, 0.34, 20, BLUE, True)
    text(s, "0                    8                   16%", 3.12, 5.53, 3.12, 0.24, 12, MUTED)
    text(s, "0                  20                  40%", 8.2, 5.53, 2.9, 0.24, 12, MUTED)
    text(s, t[4], 0.55, 5.96, 6.06, 0.31, 19, WHITE, True)
    text(s, t[5], 6.95, 5.96, 5.83, 0.31, 18, WHITE, True)
    text(s, t[6], 0.55, 6.36, 12.1, 0.31, 20, AMBER)
    text(s, t[7], 0.55, 6.78, 8.8, 0.22, 14, MINT, True)
    text(s, t[8], 0.55, 0.58, 12, 0.21, 12, MUTED)


def demo_slide(s, t):
    box(s, 0.55, 2.13, 12.23, 1.16, stroke=MINT)
    text(s, "$", 0.85, 2.4, 0.5, 0.48, 31, MINT)
    text(s, t[0], 1.47, 2.37, 10.98, 0.62, 40, WHITE, True)
    for i, value in enumerate(t[1:4]):
        text(s, value, 0.8, 3.6 + i * 0.59, 11.98, 0.5, 28, MINT if i else WHITE, True)
    for i, value in enumerate(t[4:7]):
        pill(s, value, 0.55 + i * 4.15, 5.76, 3.93, AMBER if i == 1 else MINT, 16)
    text(s, t[7], 0.55, 6.36, 12.23, 0.32, 21, MUTED)
    text(s, t[8], 0.55, 6.8, 12.23, 0.22, 14, MUTED)


def risk_slide(s, t):
    box(s, 0.55, 2.44, 2.62, 2.82, stroke=VIOLET)
    text(s, t[0], 0.78, 2.63, 2.17, 0.33, 16, VIOLET, True)
    text(s, t[1], 0.78, 3.2, 2.17, 0.45, 29, WHITE, True)
    text(s, t[2], 0.78, 3.69, 2.17, 0.36, 24, WHITE, True)
    text(s, t[3], 0.78, 4.19, 2.17, 0.37, 23, MINT)
    text(s, t[4], 0.78, 4.8, 2.17, 0.27, 14, MUTED)
    arrow(s, 3.22, 3.85, 3.66, 3.85)
    box(s, 3.72, 2.14, 5.59, 3.72, stroke=MINT)
    text(s, t[5], 4.0, 2.39, 5.0, 0.53, 34, MINT, True)
    for i, value in enumerate(t[6:11]):
        text(s, value, 4.0, 3.2 + i * 0.46, 5.05, 0.36, 19, WHITE)
    arrow(s, 9.34, 3.85, 9.8, 3.85)
    text(s, t[11], 9.97, 2.52, 2.63, 0.3, 16, MUTED, True)
    node(s, t[12], 9.87, 3.13, 2.91, 0.86)
    text(s, t[13], 9.98, 4.49, 2.62, 0.4, 24, MUTED)
    text(s, t[14], 9.98, 4.99, 2.62, 0.3, 15, AMBER, True)
    text(s, t[15], 0.55, 6.1, 12.23, 0.49, 29, MINT, True)
    text(s, t[16], 0.55, 6.67, 6.55, 0.26, 14, MUTED)
    text(s, t[17], 7.05, 6.67, 5.73, 0.26, 13, MUTED)


def agentic_slide(s, t):
    for column, ink in [(0, VIOLET), (1, MINT)]:
        x = 0.55 + column * 6.3
        box(s, x, 2.14, 5.93, 3.4, stroke=LINE)
        text(s, t[column * 5], x + 0.25, 2.39, 5.4, 0.45, 18, ink, True)
        for i in range(4):
            text(s, t[column * 5 + i + 1], x + 0.25, 3.1 + i * 0.51, 5.4, 0.42, 23)
    text(s, t[10], 0.55, 5.85, 6.15, 0.59, 33, VIOLET, True)
    text(s, t[11], 6.85, 5.85, 5.93, 0.59, 33, MINT, True)
    text(s, t[12], 0.55, 6.68, 12.23, 0.25, 14, MUTED)


def roadmap_slide(s, t):
    for i in range(4):
        x = 0.55 + i * 3.15
        ink = MINT if i == 0 else AMBER if i == 1 else MUTED
        box(s, x, 2.43, 2.77, 3.15, stroke=ink)
        text(s, t[i * 4], x + 0.21, 2.69, 2.35, 0.4, 20, ink, True)
        text(s, "\n".join(t[i * 4 + 1 : i * 4 + 4]), x + 0.21, 3.39, 2.35, 1.63, 25)
        if i < 3:
            arrow(s, x + 2.81, 4.0, x + 3.1, 4.0, MUTED)
    text(s, t[16], 0.55, 6.03, 12.23, 0.55, 31, MINT, True)
    text(s, t[17], 0.55, 6.68, 12.23, 0.27, 16, MUTED)


def closing_slide(s, t):
    text(s, t[0], 0.55, 1.31, 12.23, 0.7, 30, MUTED)
    text(s, t[1], 0.55, 2.33, 12.23, 0.89, 53, WHITE, True)
    text(s, t[2], 0.55, 3.34, 12.23, 0.97, 60, MINT, True)
    for i, value in enumerate(t[3:7]):
        pill(s, value, 0.55 + i * 3.15, 5.25, 2.78, MINT, 16)
    text(s, t[7], 0.55, 6.17, 5.2, 0.43, 26, WHITE, True)
    text(s, t[8], 6.4, 6.24, 6.38, 0.34, 20, MUTED, align=PP_ALIGN.RIGHT)


def validate(path):
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    with zipfile.ZipFile(path) as archive:
        if archive.testzip() is not None:
            raise ValueError("Corrupt ZIP container")
        names = set(archive.namelist())
        slide_names = sorted(n for n in names if re.fullmatch(r"ppt/slides/slide\d+\.xml", n))
        if len(slide_names) != 11:
            raise ValueError("Incorrect slide count")
        for name in names:
            if name.endswith(".rels"):
                base = posixpath.dirname(posixpath.dirname(name))
                for relation in ElementTree.fromstring(archive.read(name)):
                    if relation.get("TargetMode") == "External":
                        raise ValueError(f"Unexpected external relationship in {name}")
                    target = relation.attrib["Target"]
                    resolved = posixpath.normpath(posixpath.join(base, target)).lstrip("/")
                    if resolved not in names:
                        raise ValueError(f"Missing relationship: {name} -> {target}")
        for name in slide_names:
            content = ElementTree.fromstring(archive.read(name))
            if len(" ".join(content.itertext()).strip()) < 50:
                raise ValueError(f"Empty slide: {name}")
            if not content.findall(".//a:t", ns):
                raise ValueError(f"Slide has no readable text: {name}")
    deck = Presentation(path)
    for number, slide in enumerate(deck.slides, 1):
        if not slide.notes_slide.notes_text_frame.text.strip():
            raise ValueError(f"Missing speaker notes: {number}")
        for shape in slide.shapes:
            if (
                shape.left < 0
                or shape.top < 0
                or shape.left + shape.width > deck.slide_width + 10
                or shape.top + shape.height > deck.slide_height + 10
            ):
                raise ValueError(f"Out-of-bounds shape: slide {number}, {shape.name}")
    return {
        "slides": len(deck.slides),
        "aspect_ratio": "16:9",
        "zip_integrity": "passed",
        "internal_relationships": "passed",
        "external_relationships": 0,
        "nonempty_slides": "passed",
        "shape_bounds": "passed",
        "speaker_notes": "11/11",
        "bytes": path.stat().st_size,
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "native_powerpoint_keynote_check": "Requires presenter machine; not available here",
    }


def main():
    prs = Presentation()
    prs.slide_width = Inches(WIDTH)
    prs.slide_height = Inches(HEIGHT)
    prs.core_properties.title = "Agentic Quant Lab"
    prs.core_properties.subject = "Evidence-first research; synthetic offline demonstration"
    prs.core_properties.author = "Agentic Quant Lab"
    prs.core_properties.keywords = "research, evidence, falsification, risk, dry run"
    script = (DEMO / "presentation-script.md").read_text()
    scripts = re.findall(
        r"^## Slide \d+ — [^\n]+\n(.*?)(?=^## Slide |^## Source and scope|\Z)",
        script,
        re.M | re.S,
    )
    if len(scripts) != 11:
        raise ValueError("Expected eleven presenter scripts")
    layouts = [
        title_slide,
        problem_slide,
        architecture_slide,
        controls_slide,
        research_slide,
        trend_slide,
        demo_slide,
        risk_slide,
        agentic_slide,
        roadmap_slide,
        closing_slide,
    ]
    sections = [
        "THE IDEA",
        "THE TRUST PROBLEM",
        "TARGET ARCHITECTURE",
        "SCIENTIFIC CONTROLS",
        "REAL RESEARCH",
        "REAL RESEARCH",
        "DEMO FIXTURE",
        "DEMO FIXTURE / SAFETY",
        "DIVISION OF RESPONSIBILITY",
        "EVIDENCE BEFORE EXECUTION",
        "THE TAKEAWAY",
    ]
    for (number, title, copy, specification), layout, section, script in zip(
        source(), layouts, sections, scripts, strict=True
    ):
        notes = script + "\n\nEDITABLE SLIDE SPECIFICATION / SOURCES\n" + specification
        layout(new_slide(prs, number, title, section, notes), copy)
    path = DEMO / "Agentic-Quant-Lab-Demo.pptx"
    prs.save(path)
    report = validate(path)
    (DEMO / "assets").mkdir(exist_ok=True)
    (DEMO / "assets/deck-validation.json").write_text(json.dumps(report, indent=2) + "\n")
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
