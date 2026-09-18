"""Generate an editable deck from the slide source, script, and real research CSV."""

import argparse
import csv
import hashlib
import io
import json
import math
import posixpath
import re
import shutil
import subprocess
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.presentation import Presentation as PresentationDocument
from pptx.shapes.autoshape import Shape
from pptx.shapes.connector import Connector
from pptx.slide import Slide
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
DEMO = ROOT / "demo"
ASSETS = DEMO / "assets"
DECK = DEMO / "Agentic-Quant-Lab-Demo.pptx"
METRICS = ROOT / "research/etf_trend/results/metrics.csv"
BG = "0B1220"
PANEL = "142137"
WHITE = "F4F7FC"
MUTED = "B1C0D4"
MINT = "6DE5C1"
VIOLET = "B9A3FF"
AMBER = "F5C978"
RED = "FF929D"
BLUE = "90C7F6"
LINE = "34465F"
FONT = "Arial"
WIDTH = 13 + 1 / 3
HEIGHT = 7.5
SLIDE_COUNT = 11
COPY_COUNTS = [9, 17, 16, 15, 19, 9, 10, 20, 17, 18, 9]
MODELS = [
    ("SPY", "SPY"),
    ("STATIC80", "Static 80/20"),
    ("STATIC70", "Static 70/30"),
    ("ABS12", "Absolute12"),
    ("SMA10", "SMA10"),
]


def color(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def box(
    slide: Slide,
    x: float,
    y: float,
    w: float,
    h: float,
    fill: str = PANEL,
    stroke: str | None = None,
    rounded: bool = True,
) -> Shape:
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        shape.adjustments[0] = 0.12
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    if stroke:
        shape.line.color.rgb = color(stroke)
        shape.line.width = Pt(1.25)
    else:
        shape.line.fill.background()
    return shape


def text(
    slide: Slide,
    value: str,
    x: float,
    y: float,
    w: float,
    h: float,
    size: float = 24,
    ink: str = WHITE,
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    *,
    metadata: bool = False,
) -> Shape:
    if size < 20 and not metadata:
        raise ValueError(f"Body text is smaller than 20 pt: {value}")
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = f"{'metadata' if metadata else 'body'}-text-{shape.shape_id}"
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0)
    frame.margin_top = frame.margin_bottom = Inches(0)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for index, line in enumerate(value.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.alignment = align
        paragraph.space_before = paragraph.space_after = Pt(0)
        paragraph.line_spacing = 1.05
        paragraph.font.name = FONT
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color(ink)
    return shape


def arrow(
    slide: Slide,
    x1: float,
    y1: float,
    x2: float,
    y2: float,
    ink: str = MINT,
    head: bool = True,
) -> Connector:
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    line.line.color.rgb = color(ink)
    line.line.width = Pt(1.7)
    if head:
        end = OxmlElement("a:tailEnd")
        end.set("type", "triangle")
        line.line._get_or_add_ln().append(end)
    return line


def pill(slide: Slide, value: str, x: float, y: float, w: float, ink: str = MINT) -> None:
    box(slide, x, y, w, 0.48, stroke=ink)
    text(slide, value, x + 0.09, y + 0.01, w - 0.18, 0.46, 20, ink, True, PP_ALIGN.CENTER)


def node(
    slide: Slide,
    value: str,
    x: float,
    y: float,
    w: float,
    h: float = 0.75,
    ink: str = MINT,
) -> Shape:
    shape = box(slide, x, y, w, h, stroke=ink)
    text(slide, value, x + 0.07, y + 0.02, w - 0.14, h - 0.04, 20, ink, True, PP_ALIGN.CENTER)
    return shape


def source() -> list[tuple[int, str, list[str], str]]:
    raw = (DEMO / "slides.md").read_text(encoding="utf-8")
    pattern = r"^## Slide (\d+) — ([^\n]+)\n(.*?)(?=^## Slide |^## Source-of-truth|\Z)"
    result = []
    for number, title, body in re.findall(pattern, raw, re.M | re.S):
        copy = re.search(r"### On-slide text\n```text\n(.*?)\n```", body, re.S)
        if not copy:
            raise ValueError(f"Missing slide text: {number}")
        for heading in ("Visual layout", "Data source", "Speaker-note summary"):
            if f"### {heading}\n" not in body:
                raise ValueError(f"Missing {heading} on slide {number}")
        result.append((int(number), title, copy[1].splitlines(), body))
    if [entry[0] for entry in result] != list(range(1, SLIDE_COUNT + 1)):
        raise ValueError("Expected exactly eleven ordered slide specifications")
    for (number, _, copy, _), expected in zip(result, COPY_COUNTS, strict=True):
        if len(copy) != expected:
            raise ValueError(f"Slide {number}: expected {expected} text lines, got {len(copy)}")
    return result


def scripts() -> list[str]:
    raw = (DEMO / "presentation-script.md").read_text(encoding="utf-8")
    pattern = r"^## Slide (\d+) — ([^\n]+)\n(.*?)(?=^## Slide |^## Source and scope|\Z)"
    sections = re.findall(pattern, raw, re.M | re.S)
    specifications = source()
    if len(sections) != SLIDE_COUNT:
        raise ValueError("Expected eleven presenter scripts")
    seconds = 0
    for (number, title, body), (expected, slide_title, _, _) in zip(
        sections, specifications, strict=True
    ):
        if int(number) != expected or title != slide_title:
            raise ValueError(f"Script title/order does not match slide {number}")
        timing = re.search(r"\*\*Timing:\*\* (\d+) seconds", body)
        if not timing:
            raise ValueError(f"Missing timing on slide {number}")
        seconds += int(timing[1])
        for heading in ("Goal", "Script", "Transition"):
            if f"### {heading}\n" not in body:
                raise ValueError(f"Missing {heading} on slide {number}")
    if seconds != 400:
        raise ValueError(f"The presentation must total 400 seconds, got {seconds}")
    return [section[2] for section in sections]


def trend_metrics() -> dict[str, dict[str, float]]:
    with METRICS.open(encoding="utf-8", newline="") as handle:
        rows = [
            row
            for row in csv.DictReader(handle)
            if row["defense"] == "BIL" and row["one_way_cost_bps"] == "20"
        ]
    result = {}
    for model, _ in MODELS:
        matching = [row for row in rows if row["model"] == model]
        if len(matching) != 1:
            raise ValueError(f"Expected exactly one committed BIL/20bps row for {model}")
        cagr = float(matching[0]["cagr"]) * 100
        drawdown = float(matching[0]["max_drawdown"]) * 100
        if not (math.isfinite(cagr) and 0 < cagr <= 16):
            raise ValueError(f"CAGR outside this chart's 0-16% scale: {model}")
        if not (math.isfinite(drawdown) and -40 <= drawdown <= 0):
            raise ValueError(f"Drawdown outside this chart's 0-40% magnitude scale: {model}")
        result[model] = {"cagr_pct": cagr, "max_drawdown_pct": drawdown}
    return result


def new_slide(prs: PresentationDocument, number: int, title: str, section: str, note: str) -> Slide:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color(BG)
    text(slide, "AQL / " + section, 0.55, 0.27, 11.8, 0.3, 14, MINT, True, metadata=True)
    if number not in (1, 11):
        text(slide, title, 0.55, 0.8, 12.23, 1.12, 34, bold=True)
    text(slide, "AGENTIC QUANT LAB", 0.55, 7.07, 7, 0.22, 12, MUTED, metadata=True)
    text(
        slide,
        f"{number:02} / {SLIDE_COUNT}",
        11.7,
        7.05,
        1.08,
        0.25,
        12,
        MUTED,
        align=PP_ALIGN.RIGHT,
        metadata=True,
    )
    for index in range(SLIDE_COUNT):
        box(
            slide,
            0.55 + index * 1.115,
            7.37,
            1.045,
            0.025,
            fill=MINT if index < number else LINE,
            rounded=False,
        )
    notes_frame = slide.notes_slide.notes_text_frame
    if notes_frame is None:
        raise ValueError(f"Slide {number} has no notes placeholder")
    notes_frame.text = note
    return slide


def title_slide(s: Slide, t: list[str]) -> None:
    text(s, "Agentic\nQuant Lab", 0.55, 1.25, 7.6, 2.0, 62, bold=True)
    text(s, "\n".join(t[:2]), 0.6, 3.52, 7.0, 1.1, 26, MUTED)
    text(s, "\n".join(t[2:4]), 0.6, 5.25, 7.0, 0.94, 28, MINT, True)
    for i, value in enumerate(t[4:8]):
        y = 1.5 + i * 1.14
        node(s, value, 8.55, y, 3.7, 0.83, VIOLET if i == 0 else MINT)
        if i < 3:
            arrow(s, 10.4, y + 0.85, 10.4, y + 1.1)
    text(s, t[8], 8.05, 6.32, 4.5, 0.36, 16, MUTED, align=PP_ALIGN.CENTER, metadata=True)


def problem_slide(s: Slide, t: list[str]) -> None:
    box(s, 0.55, 2.05, 12.23, 1.25, fill="251D2C")
    text(s, t[0], 0.8, 2.18, 3.0, 0.8, 20, RED, True)
    for i, value in enumerate(t[1:4]):
        x = 4.05 + i * 2.9
        node(s, value, x, 2.33, 2.43, 0.66, RED)
        if i < 2:
            arrow(s, x + 2.46, 2.66, x + 2.83, 2.66, RED)
    box(s, 0.55, 3.52, 12.23, 1.54)
    text(s, t[4], 0.8, 3.65, 4.2, 0.36, 20, MINT, True)
    text(s, t[16], 6.0, 3.68, 6.5, 0.3, 16, MUTED, metadata=True)
    for i, value in enumerate(t[5:10]):
        x = 0.8 + i * 2.4
        node(s, value, x, 4.18, 2.16, 0.64)
        if i < 4:
            arrow(s, x + 2.17, 4.5, x + 2.35, 4.5)
    for i, value in enumerate(t[10:14]):
        text(s, value, 0.6 + i * 3.1, 5.22, 2.9, 0.62, 20, MUTED)
    text(s, t[14], 0.55, 6.0, 12.23, 0.4, 28, WHITE, True)
    text(s, t[15], 0.55, 6.44, 12.23, 0.46, 32, MINT, True)


def architecture_slide(s: Slide, t: list[str]) -> None:
    pill(s, t[0], 0.55, 2.0, 2.75, VIOLET)
    pill(s, t[1], 3.5, 2.0, 5.25)
    node(s, t[2], 0.55, 2.75, 2.5, ink=VIOLET)
    node(s, t[3], 0.55, 4.02, 2.5, ink=VIOLET)
    arrow(s, 1.8, 3.53, 1.8, 3.98, VIOLET)
    arrow(s, 3.06, 4.39, 3.28, 4.39, VIOLET, False)
    arrow(s, 3.28, 4.39, 3.28, 3.13, VIOLET, False)
    arrow(s, 3.28, 3.13, 3.49, 3.13, VIOLET)
    xs = [3.5, 5.85, 8.2, 10.55]
    for i, x in enumerate(xs):
        node(s, t[4 + i], x, 2.75, 2.18)
        if i < 3:
            arrow(s, x + 2.2, 3.13, xs[i + 1] - 0.02, 3.13)
    arrow(s, 11.64, 3.53, 11.64, 4.22)
    for i, x in enumerate(reversed(xs)):
        shape = node(s, t[8 + i], x, 4.25, 2.18, ink=MUTED if i == 3 else MINT)
        if i == 3:
            shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        if i < 3:
            arrow(s, x - 0.01, 4.63, x - 0.16, 4.63, MUTED if i == 2 else MINT)
    text(s, t[12], 3.5, 5.11, 2.18, 0.3, 16, MUTED, align=PP_ALIGN.CENTER, metadata=True)
    box(s, 0.55, 5.78, 12.18, 0.83, fill="182435", stroke=LINE)
    text(s, t[13], 0.8, 5.87, 3.6, 0.68, 20, BLUE, True)
    text(s, t[14], 4.55, 5.94, 7.9, 0.43, 21, MUTED)
    arrow(s, 1.8, 4.81, 1.8, 5.75, BLUE)
    arrow(s, 7.83, 3.53, 8.11, 3.72, BLUE, False)
    arrow(s, 8.11, 3.72, 8.11, 5.75, BLUE)
    arrow(s, 9.29, 5.03, 9.29, 5.75, BLUE)
    text(s, t[15], 0.55, 6.72, 12.2, 0.27, 14, MUTED, metadata=True)


def controls_slide(s: Slide, t: list[str]) -> None:
    for i, (symbol, ink) in enumerate([("0", BLUE), ("+", MINT), ("t+1", AMBER)]):
        x = 0.55 + i * 4.15
        box(s, x, 2.14, 3.93, 3.05, stroke=LINE)
        text(s, t[i * 3], x + 0.24, 2.39, 3.45, 0.4, 20, ink, True)
        text(s, symbol, x + 0.24, 2.99, 3.45, 0.88, 56, ink, True)
        text(s, "\n".join(t[i * 3 + 1 : i * 3 + 3]), x + 0.24, 4.01, 3.45, 0.87, 26)
    for i, value in enumerate(t[9:13]):
        text(s, value, 0.55 + i * 3.1, 5.4, 2.93, 0.65, 20, MUTED)
    text(s, t[13], 0.55, 6.16, 12.23, 0.5, 29, MINT, True)
    text(s, t[14], 0.55, 6.78, 12.23, 0.27, 14, MUTED, metadata=True)


def research_slide(s: Slide, t: list[str]) -> None:
    text(s, t[0], 10.0, 0.27, 2.77, 0.3, 14, AMBER, True, PP_ALIGN.RIGHT, metadata=True)
    for i in range(5):
        name, status, reason = t[1 + i * 3 : 4 + i * 3]
        y = 2.05 + i * 0.77
        ink = AMBER if i in (0, 3) else BLUE
        box(s, 0.55, y, 12.23, 0.66)
        box(s, 0.55, y, 0.055, 0.66, fill=ink, rounded=False)
        text(s, name, 0.77, y + 0.07, 3.7, 0.51, 20, bold=True)
        text(s, status, 4.6, y + 0.07, 2.73, 0.51, 20, ink, True)
        text(s, reason, 7.42, y + 0.06, 5.1, 0.53, 20, MUTED)
    text(s, t[16], 0.55, 6.15, 4.6, 0.5, 30, MINT, True)
    text(s, t[17], 5.15, 6.16, 7.6, 0.48, 25, WHITE, True)
    text(s, t[18], 0.55, 6.78, 12.23, 0.27, 16, MUTED, metadata=True)


def trend_slide(s: Slide, t: list[str]) -> None:
    text(s, t[0], 0.55, 2.0, 12, 0.35, 20, AMBER, True)
    text(s, t[1], 0.55, 2.43, 12, 0.31, 18, MUTED, metadata=True)
    text(s, t[2], 3.12, 2.89, 3.5, 0.36, 20, MINT, True)
    text(s, t[3], 8.12, 2.89, 4.65, 0.36, 20, BLUE, True)
    rows = trend_metrics()
    for i, (model, label) in enumerate(MODELS):
        y = 3.4 + i * 0.43
        cagr, drawdown = rows[model]["cagr_pct"], rows[model]["max_drawdown_pct"]
        ink = AMBER if model in ("ABS12", "SMA10") else MINT
        text(s, label, 0.55, y - 0.025, 2.45, 0.36, 20, bold=True)
        box(s, 3.12, y + 0.035, cagr / 16 * 3.0, 0.21, fill=ink, rounded=False)
        text(s, f"{cagr:.2f}%", 6.26, y - 0.025, 1.45, 0.36, 20, ink, True)
        box(s, 8.12, y + 0.035, abs(drawdown) / 40 * 2.72, 0.21, fill=BLUE, rounded=False)
        text(s, f"{drawdown:.2f}%", 11.04, y - 0.025, 1.74, 0.36, 20, BLUE, True)
    for origin, width, maximum in [(3.12, 3.0, 16), (8.12, 2.72, 40)]:
        for fraction in (0, 0.5, 1):
            label = f"{maximum * fraction:g}" + ("%" if fraction == 1 else "")
            text(
                s,
                label,
                origin + width * fraction - 0.08,
                5.56,
                0.55,
                0.25,
                14,
                MUTED,
                metadata=True,
            )
    text(s, t[4], 0.55, 5.94, 5.95, 0.69, 20, WHITE, True)
    text(s, t[5], 6.9, 5.94, 5.87, 0.69, 20, WHITE, True)
    text(s, t[6], 0.55, 6.65, 12.2, 0.35, 20, AMBER)
    text(s, t[7], 0.55, 0.57, 12, 0.36, 20, MINT, True)
    qualifier, limitation = t[8].split(" · ", 1)
    text(s, qualifier, 8.7, 2.0, 4.08, 0.35, 20, AMBER, True)
    text(s, limitation, 4.75, 7.06, 6.7, 0.23, 12, MUTED, metadata=True)


def demo_slide(s: Slide, t: list[str]) -> None:
    box(s, 0.55, 2.13, 12.23, 1.13, stroke=MINT)
    text(s, "$", 0.85, 2.4, 0.5, 0.48, 31, MINT)
    text(s, t[0], 1.47, 2.37, 10.98, 0.62, 40, WHITE, True)
    for i, value in enumerate(t[1:4]):
        text(s, value, 0.8, 3.46 + i * 0.59, 11.98, 0.5, 28, MINT if i else WHITE, True)
    for i, value in enumerate(t[4:8]):
        x = 0.55 + i * 3.15
        ink = AMBER if i == 0 else MINT
        box(s, x, 5.4, 2.78, 0.82, stroke=ink)
        text(
            s,
            value.replace(": ", ":\n"),
            x + 0.08,
            5.44,
            2.62,
            0.74,
            20,
            ink,
            True,
            PP_ALIGN.CENTER,
        )
    text(s, t[8], 0.55, 6.4, 12.23, 0.35, 20, MUTED)
    text(s, t[9], 0.55, 6.83, 12.23, 0.23, 14, MUTED, metadata=True)


def risk_slide(s: Slide, t: list[str]) -> None:
    box(s, 0.55, 2.64, 2.62, 2.82, stroke=VIOLET)
    text(s, t[0], 0.75, 2.83, 2.22, 0.4, 20, VIOLET, True)
    text(s, t[1], 0.78, 3.37, 2.17, 0.46, 29, WHITE, True)
    text(s, t[2], 0.78, 3.9, 2.17, 0.4, 24, WHITE, True)
    text(s, t[3], 0.78, 4.44, 2.17, 0.4, 23, MINT)
    text(s, t[4], 0.78, 5.06, 2.17, 0.27, 14, MUTED, metadata=True)
    arrow(s, 3.22, 4.05, 3.66, 4.05)
    box(s, 3.72, 2.14, 5.59, 3.9, stroke=MINT)
    text(s, t[5], 4.0, 2.39, 5.0, 0.53, 34, MINT, True)
    for i, value in enumerate(t[6:13]):
        text(s, value, 3.97, 3.12 + i * 0.39, 5.1, 0.37, 20, WHITE)
    arrow(s, 9.34, 4.05, 9.8, 4.05)
    text(s, t[13], 9.97, 2.72, 2.63, 0.34, 20, MUTED, True)
    node(s, t[14].replace(" ADAPTER", "\nADAPTER"), 9.87, 3.42, 2.91, 1.0)
    text(s, t[15], 9.98, 4.93, 2.62, 0.42, 24, MUTED)
    text(s, t[16], 9.98, 5.46, 2.62, 0.35, 20, AMBER, True)
    text(s, t[17], 0.55, 6.25, 12.23, 0.49, 29, MINT, True)
    text(s, t[18], 0.55, 6.81, 7.0, 0.24, 14, MUTED, metadata=True)
    text(s, t[19], 7.6, 6.81, 5.18, 0.24, 12, MUTED, metadata=True)


def agentic_slide(s: Slide, t: list[str]) -> None:
    for column, ink in [(0, VIOLET), (1, MINT)]:
        x = 0.55 + column * 6.3
        box(s, x, 2.14, 5.93, 3.63, stroke=LINE)
        text(s, t[column * 7], x + 0.25, 2.39, 5.4, 0.45, 20, ink, True)
        for i in range(6):
            text(s, t[column * 7 + i + 1], x + 0.25, 3.04 + i * 0.41, 5.4, 0.39, 23)
    text(s, t[14], 0.55, 6.03, 6.15, 0.58, 33, VIOLET, True)
    text(s, t[15], 6.85, 6.03, 5.93, 0.58, 33, MINT, True)
    text(s, t[16], 0.55, 6.79, 12.23, 0.26, 14, MUTED, metadata=True)


def roadmap_slide(s: Slide, t: list[str]) -> None:
    for i in range(4):
        x = 0.55 + i * 3.15
        ink = MINT if i == 0 else AMBER if i == 1 else MUTED
        box(s, x, 2.43, 2.77, 3.15, stroke=ink)
        text(s, t[i * 4], x + 0.21, 2.69, 2.35, 0.4, 20, ink, True)
        text(s, "\n".join(t[i * 4 + 1 : i * 4 + 4]), x + 0.21, 3.33, 2.35, 1.83, 25)
        if i < 3:
            arrow(s, x + 2.81, 4.0, x + 3.1, 4.0, MUTED)
    text(s, t[16], 0.55, 6.03, 12.23, 0.55, 31, MINT, True)
    text(s, t[17], 0.55, 6.78, 12.23, 0.27, 16, MUTED, metadata=True)


def closing_slide(s: Slide, t: list[str]) -> None:
    text(s, t[0], 0.55, 1.31, 12.23, 0.7, 30, MUTED)
    text(s, t[1], 0.55, 2.33, 12.23, 0.89, 53, WHITE, True)
    text(s, t[2], 0.55, 3.34, 12.23, 0.97, 60, MINT, True)
    for i, value in enumerate(t[3:7]):
        pill(s, value, 0.55 + i * 3.15, 5.25, 2.78)
    text(s, t[7], 0.55, 6.17, 5.2, 0.43, 26, WHITE, True)
    text(s, t[8], 6.4, 6.24, 6.38, 0.34, 20, MUTED, align=PP_ALIGN.RIGHT)


def save_reproducibly(prs: PresentationDocument, path: Path) -> None:
    data = io.BytesIO()
    prs.save(data)
    with zipfile.ZipFile(data) as source_archive, zipfile.ZipFile(path, "w") as target:
        for name in sorted(source_archive.namelist()):
            info = zipfile.ZipInfo(name, date_time=(2026, 9, 18, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(info, source_archive.read(name))


def validate(path: Path) -> dict[str, object]:
    drawing_ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    relationship_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    image_count = 0
    with zipfile.ZipFile(path) as archive:
        if archive.testzip() is not None:
            raise ValueError("Corrupt ZIP container")
        names = set(archive.namelist())
        if len(names) != len(archive.namelist()):
            raise ValueError("Duplicate ZIP entries")
        slide_names = sorted(n for n in names if re.fullmatch(r"ppt/slides/slide\d+\.xml", n))
        if len(slide_names) != SLIDE_COUNT:
            raise ValueError("Incorrect slide count")
        if "[Content_Types].xml" not in names:
            raise ValueError("Missing content types")
        for name in sorted(names):
            if not name.endswith((".xml", ".rels")):
                continue
            content = ElementTree.fromstring(archive.read(name))
            if name.endswith(".rels"):
                base = posixpath.dirname(posixpath.dirname(name))
                for relation in content:
                    if relation.get("TargetMode") == "External":
                        raise ValueError(f"Unexpected external relationship in {name}")
                    target = relation.attrib["Target"]
                    resolved = posixpath.normpath(posixpath.join(base, target)).lstrip("/")
                    if resolved not in names:
                        raise ValueError(f"Missing relationship: {name} -> {target}")
            else:
                ids = {
                    value
                    for element in content.iter()
                    for attribute, value in element.attrib.items()
                    if attribute
                    in {relationship_ns + "id", relationship_ns + "embed", relationship_ns + "link"}
                }
                if ids:
                    relations_name = posixpath.join(
                        posixpath.dirname(name), "_rels", posixpath.basename(name) + ".rels"
                    )
                    if relations_name not in names:
                        raise ValueError(f"Missing relationships file for {name}")
                    defined = {
                        relation.attrib["Id"]
                        for relation in ElementTree.fromstring(archive.read(relations_name))
                    }
                    if ids - defined:
                        raise ValueError(f"Unresolved relationship IDs in {name}: {ids - defined}")
            if name in slide_names and not content.findall(".//a:t", drawing_ns):
                raise ValueError(f"Slide has no readable text: {name}")
        for name in names:
            if name.startswith("ppt/media/"):
                with Image.open(io.BytesIO(archive.read(name))) as image:
                    image.verify()
                image_count += 1
    deck = Presentation(str(path))
    width, height = deck.slide_width, deck.slide_height
    if width is None or height is None or height <= 0:
        raise ValueError("Deck has no valid slide dimensions")
    if abs(width / height - 16 / 9) > 0.000001:
        raise ValueError("Deck is not 16:9")
    minimum = 100.0
    for number, slide in enumerate(deck.slides, 1):
        notes_frame = slide.notes_slide.notes_text_frame
        if notes_frame is None:
            raise ValueError(f"Slide {number} has no notes placeholder")
        notes = notes_frame.text
        if not all(part in notes for part in ("### Script", "### Transition", "### Data source")):
            raise ValueError(f"Incomplete speaker notes/source: {number}")
        for shape in slide.shapes:
            if (
                shape.left < 0
                or shape.top < 0
                or shape.left + shape.width > width + 10
                or shape.top + shape.height > height + 10
            ):
                raise ValueError(f"Out-of-bounds shape: slide {number}, {shape.name}")
            if isinstance(shape, Shape) and shape.name.startswith("body-text-"):
                for paragraph in shape.text_frame.paragraphs:
                    size = paragraph.font.size
                    if size is None or size.pt < 20:
                        raise ValueError(f"Body text below 20 pt on slide {number}")
                    minimum = min(minimum, size.pt)
    return {
        "slides": len(deck.slides),
        "aspect_ratio": "16:9",
        "zip_integrity": "passed",
        "all_xml_parseable": "passed",
        "internal_relationships": "passed",
        "referenced_relationship_ids": "passed",
        "external_relationships": 0,
        "embedded_images": image_count,
        "visuals": "Editable native shapes; no linked images or external assets",
        "shape_bounds": "passed",
        "minimum_body_font_pt": minimum,
        "small_type_scope": "Metadata, sources, chart ticks, and footers only",
        "speaker_notes": "11/11 with script, transition, and source",
        "presentation_seconds": 400,
        "metric_source": str(METRICS.relative_to(ROOT)),
        "metric_source_sha256": hashlib.sha256(METRICS.read_bytes()).hexdigest(),
        "metric_filter": {"defense": "BIL", "one_way_cost_bps": 20},
        "real_research_metrics": trend_metrics(),
        "bytes": path.stat().st_size,
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "native_powerpoint_keynote_check": "Requires presenter machine; not available here",
    }


def render(path: Path) -> dict[str, object]:
    office = shutil.which("libreoffice") or shutil.which("soffice")
    rasterizer = shutil.which("pdftoppm")
    if not office or not rasterizer:
        raise RuntimeError(
            "Local rendering requires LibreOffice and pdftoppm. "
            "The PPTX has already been generated and structurally validated; "
            "do not discard it or claim stale previews are current."
        )
    with tempfile.TemporaryDirectory(prefix="aql-deck-render-") as temporary:
        directory = Path(temporary)
        subprocess.run(
            [
                office,
                f"-env:UserInstallation={(directory / 'profile').as_uri()}",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(directory),
                str(path),
            ],
            check=True,
            timeout=180,
        )
        pdf = directory / path.with_suffix(".pdf").name
        if not pdf.is_file() or not pdf.read_bytes().startswith(b"%PDF-"):
            raise ValueError("LibreOffice did not produce a valid PDF header")
        subprocess.run(
            [rasterizer, "-scale-to", "1600", "-png", str(pdf), str(directory / "slide")],
            check=True,
            timeout=180,
        )
        pages = sorted(
            directory.glob("slide-*.png"),
            key=lambda entry: int(entry.stem.split("-")[-1]),
        )
        if len(pages) != SLIDE_COUNT:
            raise ValueError(f"Expected eleven rendered pages, got {len(pages)}")
        slides_directory = ASSETS / "slides"
        slides_directory.mkdir(exist_ok=True)
        sheet = Image.new("RGB", (1992, 1626), "#" + BG)
        draw = ImageDraw.Draw(sheet)
        label_font = ImageFont.load_default(size=16)
        for index, page in enumerate(pages):
            with Image.open(page) as image:
                image.load()
                if image.width / image.height != 16 / 9:
                    raise ValueError(f"Unexpected rendered aspect ratio on slide {index + 1}")
                preview = image.convert("RGB").resize((640, 360), Image.Resampling.LANCZOS)
            x, y = 18 + (index % 3) * 658, 18 + (index // 3) * 402
            sheet.paste(preview, (x, y))
            draw.text((x, y + 367), f"SLIDE {index + 1:02}", font=label_font, fill="#" + MUTED)
            shutil.copyfile(page, slides_directory / f"slide-{index + 1:02}.png")
        sheet.save(ASSETS / "deck-preview.png", optimize=True)
        shutil.copyfile(pdf, path.with_suffix(".pdf"))
    return {
        "renderer": "LibreOffice / Poppler, local",
        "rendered_pages": SLIDE_COUNT,
        "slide_images": "demo/assets/slides/slide-01.png through slide-11.png",
        "slide_image_sha256": {
            f"slide-{number:02}.png": hashlib.sha256(
                (ASSETS / f"slides/slide-{number:02}.png").read_bytes()
            ).hexdigest()
            for number in range(1, SLIDE_COUNT + 1)
        },
        "pdf_sha256": hashlib.sha256(path.with_suffix(".pdf").read_bytes()).hexdigest(),
        "preview_sha256": hashlib.sha256((ASSETS / "deck-preview.png").read_bytes()).hexdigest(),
        "visual_inspection": "Pending inspection of all eleven rendered slides",
    }


def matching_rendering(
    report: dict[str, object], validation_path: Path
) -> dict[str, object] | None:
    if not validation_path.exists():
        return None
    previous = json.loads(validation_path.read_text(encoding="utf-8"))
    if not isinstance(previous, dict):
        raise ValueError("Existing deck validation report is not a JSON object")
    rendering = previous.get("rendering")
    if previous.get("sha256") != report["sha256"] or not isinstance(rendering, dict):
        return None
    files = {
        DECK.with_suffix(".pdf"): rendering.get("pdf_sha256"),
        ASSETS / "deck-preview.png": rendering.get("preview_sha256"),
    }
    image_hashes = rendering.get("slide_image_sha256")
    expected_images = {f"slide-{number:02}.png" for number in range(1, SLIDE_COUNT + 1)}
    if not isinstance(image_hashes, dict) or set(image_hashes) != expected_images:
        return None
    files.update({ASSETS / "slides" / name: image_hashes[name] for name in expected_images})
    if any(
        not path.is_file() or hashlib.sha256(path.read_bytes()).hexdigest() != expected
        for path, expected in files.items()
    ):
        return None
    return rendering


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--render", action="store_true", help="Refresh PDF and PNGs with local tools"
    )
    parser.add_argument("--validate-only", action="store_true", help="Check the existing PPTX")
    args = parser.parse_args()
    if args.validate_only:
        print(json.dumps(validate(DECK), indent=2))
        return
    prs = Presentation()
    prs.slide_width = Inches(WIDTH)
    prs.slide_height = Inches(HEIGHT)
    prs.core_properties.title = "Agentic Quant Lab"
    prs.core_properties.subject = "Evidence-first research; synthetic offline demonstration"
    prs.core_properties.author = "Agentic Quant Lab"
    prs.core_properties.keywords = "research, evidence, falsification, risk, dry run"
    prs.core_properties.created = prs.core_properties.modified = datetime(2026, 9, 18)
    prs.core_properties.revision = 1
    layouts = [
        title_slide,
        problem_slide,
        architecture_slide,
        controls_slide,
        research_slide,
        trend_slide,
        demo_slide,
        risk_slide,
        agentic_slide,
        roadmap_slide,
        closing_slide,
    ]
    sections = [
        "THE IDEA",
        "THE TRUST PROBLEM",
        "TARGET ARCHITECTURE",
        "SCIENTIFIC CONTROLS",
        "REAL RESEARCH",
        "REAL RESEARCH",
        "SYNTHETIC DEMO",
        "SYNTHETIC DEMO / SAFETY",
        "DIVISION OF RESPONSIBILITY",
        "EVIDENCE BEFORE EXECUTION",
        "THE TAKEAWAY",
    ]
    for (number, title, copy, specification), layout, section, script in zip(
        source(), layouts, sections, scripts(), strict=True
    ):
        notes = script + "\n\nEDITABLE SLIDE SPECIFICATION / SOURCES\n" + specification
        layout(new_slide(prs, number, title, section, notes), copy)
    save_reproducibly(prs, DECK)
    report = validate(DECK)
    ASSETS.mkdir(exist_ok=True)
    validation_path = ASSETS / "deck-validation.json"
    report["rendering"] = (
        matching_rendering(report, validation_path)
        or "Not run; any prior PDF/preview must be regenerated for this PPTX"
    )
    validation_path.write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    if args.render:
        report["rendering"] = render(DECK)
        validation_path.write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
